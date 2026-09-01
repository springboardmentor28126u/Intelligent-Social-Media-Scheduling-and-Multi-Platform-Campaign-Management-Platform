import os
from datetime import datetime, timezone
from typing import Dict, Any, List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_pdf_report(data: Dict[str, Any], output_path: str) -> str:
    """Generates a styled PDF analytics report."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Typography Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#1e293b')
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#64748b')
    )
    
    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#2563eb'),
        spaceBefore=14,
        spaceAfter=8
    )
    
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#334155')
    )
    
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#0f172a')
    )
    
    elements = []
    
    # 1. Header Banner
    title_text = data.get("title", "SocialPilot Analytics & Performance Report")
    elements.append(Paragraph(f"<b>SocialPilot</b> | {title_text}", title_style))
    generated_at = datetime.now(timezone.utc).strftime("%B %d, %Y - %H:%M UTC")
    author_name = data.get("user_name", "Marketing Team")
    elements.append(Paragraph(f"Generated for: <b>{author_name}</b> &nbsp;|&nbsp; Date: {generated_at}", subtitle_style))
    elements.append(Spacer(1, 12))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=14))
    
    # 2. Executive Summary KPI Cards Table
    summary = data.get("summary", {})
    kpi_data = [
        [
            Paragraph("<b>Total Reach</b>", cell_bold),
            Paragraph("<b>Total Engagement</b>", cell_bold),
            Paragraph("<b>Total Impressions</b>", cell_bold),
            Paragraph("<b>Followers Growth</b>", cell_bold)
        ],
        [
            Paragraph(f"<font size=14 color='#2563eb'><b>{summary.get('total_reach', '124,500'):,}</b></font>", cell_style),
            Paragraph(f"<font size=14 color='#16a34a'><b>{summary.get('total_engagement', '18,420'):,}</b></font>", cell_style),
            Paragraph(f"<font size=14 color='#9333ea'><b>{summary.get('total_impressions', '384,100'):,}</b></font>", cell_style),
            Paragraph(f"<font size=14 color='#ea580c'><b>+{summary.get('followers_growth', '2,450'):,}</b></font>", cell_style)
        ]
    ]
    
    kpi_table = Table(kpi_data, colWidths=[130, 130, 130, 130])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('PADDING', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 16))
    
    # 3. Platform Breakdown Section
    elements.append(Paragraph("Platform Performance Breakdown", section_heading))
    platforms = data.get("platforms", [
        {"platform": "Instagram", "posts": 28, "reach": "64,200", "engagement": "8,920", "rate": "13.9%"},
        {"platform": "Facebook", "posts": 24, "reach": "38,100", "engagement": "4,150", "rate": "10.8%"},
        {"platform": "LinkedIn", "posts": 16, "reach": "14,800", "engagement": "3,420", "rate": "23.1%"},
        {"platform": "Twitter / X", "posts": 42, "reach": "7,400", "engagement": "1,930", "rate": "26.0%"}
    ])
    
    plat_table_data = [
        [
            Paragraph("<b>Platform</b>", cell_bold),
            Paragraph("<b>Published Posts</b>", cell_bold),
            Paragraph("<b>Audience Reach</b>", cell_bold),
            Paragraph("<b>Total Engagement</b>", cell_bold),
            Paragraph("<b>Engagement Rate</b>", cell_bold)
        ]
    ]
    
    for p in platforms:
        plat_table_data.append([
            Paragraph(f"<b>{p.get('platform', '')}</b>", cell_style),
            Paragraph(str(p.get("posts", 0)), cell_style),
            Paragraph(str(p.get("reach", "0")), cell_style),
            Paragraph(str(p.get("engagement", "0")), cell_style),
            Paragraph(str(p.get("rate", "0%")), cell_style),
        ])
        
    plat_table = Table(plat_table_data, colWidths=[130, 95, 100, 105, 90])
    plat_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(plat_table)
    elements.append(Spacer(1, 16))
    
    # 4. Top Performing Content Section
    elements.append(Paragraph("Top Performing Published Posts", section_heading))
    posts = data.get("top_posts", [
        {"title": "Product 2.0 Launch Announcement", "platform": "Instagram", "likes": "1,840", "comments": "210", "shares": "95", "reach": "32,400"},
        {"title": "Quarterly Industry Insights Guide", "platform": "LinkedIn", "likes": "620", "comments": "88", "shares": "142", "reach": "12,100"},
        {"title": "Behind the Scenes with Engineering Team", "platform": "Facebook", "likes": "450", "comments": "54", "shares": "28", "reach": "9,800"}
    ])
    
    posts_table_data = [
        [
            Paragraph("<b>Post Title / Topic</b>", cell_bold),
            Paragraph("<b>Platform</b>", cell_bold),
            Paragraph("<b>Likes</b>", cell_bold),
            Paragraph("<b>Comments</b>", cell_bold),
            Paragraph("<b>Shares</b>", cell_bold),
            Paragraph("<b>Reach</b>", cell_bold)
        ]
    ]
    
    for p in posts:
        posts_table_data.append([
            Paragraph(p.get("title", "")[:40], cell_style),
            Paragraph(p.get("platform", ""), cell_style),
            Paragraph(str(p.get("likes", 0)), cell_style),
            Paragraph(str(p.get("comments", 0)), cell_style),
            Paragraph(str(p.get("shares", 0)), cell_style),
            Paragraph(str(p.get("reach", 0)), cell_style),
        ])
        
    posts_table = Table(posts_table_data, colWidths=[170, 80, 65, 65, 65, 75])
    posts_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(posts_table)
    elements.append(Spacer(1, 20))
    
    # 5. Footer note
    elements.append(Paragraph("<font size=8 color='#94a3b8'>Confidential - Generated by SocialPilot Platform - All rights reserved</font>", subtitle_style))
    
    doc.build(elements)
    return output_path


def generate_excel_report(data: Dict[str, Any], output_path: str) -> str:
    """Generates a multi-sheet formatted Excel workbook."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    wb = openpyxl.Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # Sheet 1: Executive Summary
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True
    
    ws_summary["A1"] = "SocialPilot - Campaign & Analytics Report"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Generated at: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws_summary["A2"].font = regular_font
    
    summary = data.get("summary", {})
    metrics = [
        ("Total Reach", summary.get("total_reach", 124500)),
        ("Total Impressions", summary.get("total_impressions", 384100)),
        ("Total Engagement", summary.get("total_engagement", 18420)),
        ("Followers Growth", summary.get("followers_growth", 2450)),
        ("Total Posts Published", summary.get("total_posts", 110)),
        ("Active Campaigns", summary.get("active_campaigns", 5))
    ]
    
    ws_summary["A4"] = "Key Performance Metric"
    ws_summary["B4"] = "Value"
    ws_summary["A4"].font = header_font
    ws_summary["A4"].fill = header_fill
    ws_summary["B4"].font = header_font
    ws_summary["B4"].fill = header_fill
    
    for row_idx, (metric_name, val) in enumerate(metrics, start=5):
        ws_summary[f"A{row_idx}"] = metric_name
        ws_summary[f"B{row_idx}"] = val
        ws_summary[f"A{row_idx}"].font = bold_font
        ws_summary[f"B{row_idx}"].font = regular_font
        ws_summary[f"A{row_idx}"].border = thin_border
        ws_summary[f"B{row_idx}"].border = thin_border
        
    # Sheet 2: Posts Performance
    ws_posts = wb.create_sheet(title="Posts Performance")
    ws_posts.views.sheetView[0].showGridLines = True
    
    post_headers = ["ID", "Title / Content", "Platform", "Status", "Scheduled At", "Published At", "Likes", "Comments", "Shares", "Reach"]
    for col_idx, h in enumerate(post_headers, start=1):
        cell = ws_posts.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    raw_posts = data.get("raw_posts", [])
    for row_idx, p in enumerate(raw_posts, start=2):
        ws_posts.cell(row=row_idx, column=1, value=p.get("id", row_idx-1)).border = thin_border
        ws_posts.cell(row=row_idx, column=2, value=p.get("content", "")[:60]).border = thin_border
        ws_posts.cell(row=row_idx, column=3, value=p.get("platform", "")).border = thin_border
        ws_posts.cell(row=row_idx, column=4, value=p.get("status", "")).border = thin_border
        ws_posts.cell(row=row_idx, column=5, value=str(p.get("scheduled_at", ""))).border = thin_border
        ws_posts.cell(row=row_idx, column=6, value=str(p.get("published_at", ""))).border = thin_border
        ws_posts.cell(row=row_idx, column=7, value=p.get("likes_count", 0)).border = thin_border
        ws_posts.cell(row=row_idx, column=8, value=p.get("comments_count", 0)).border = thin_border
        ws_posts.cell(row=row_idx, column=9, value=p.get("shares_count", 0)).border = thin_border
        ws_posts.cell(row=row_idx, column=10, value=p.get("reach_count", 0)).border = thin_border

    # Sheet 3: Campaigns ROI
    ws_camps = wb.create_sheet(title="Campaigns")
    ws_camps.views.sheetView[0].showGridLines = True
    
    camp_headers = ["Campaign Name", "Platform", "Status", "Budget ($)", "Spent ($)", "Target Reach", "Actual Reach", "ROI (x)"]
    for col_idx, h in enumerate(camp_headers, start=1):
        cell = ws_camps.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    raw_campaigns = data.get("raw_campaigns", [])
    for row_idx, c in enumerate(raw_campaigns, start=2):
        ws_camps.cell(row=row_idx, column=1, value=c.get("name", "")).border = thin_border
        ws_camps.cell(row=row_idx, column=2, value=c.get("platform", "")).border = thin_border
        ws_camps.cell(row=row_idx, column=3, value=c.get("status", "")).border = thin_border
        ws_camps.cell(row=row_idx, column=4, value=c.get("budget", 0)).border = thin_border
        ws_camps.cell(row=row_idx, column=5, value=c.get("spent", 0)).border = thin_border
        ws_camps.cell(row=row_idx, column=6, value=c.get("target_reach", 0)).border = thin_border
        ws_camps.cell(row=row_idx, column=7, value=c.get("actual_reach", 0)).border = thin_border
        ws_camps.cell(row=row_idx, column=8, value=c.get("roi", 0.0)).border = thin_border

    # Auto-adjust column widths across all sheets
    for ws in [ws_summary, ws_posts, ws_camps]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    wb.save(output_path)
    return output_path
