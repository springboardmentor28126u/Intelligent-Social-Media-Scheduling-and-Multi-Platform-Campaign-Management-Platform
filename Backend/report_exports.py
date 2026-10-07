"""Native Excel and paginated PDF reports from authorized report data."""
from datetime import datetime, timezone
from io import BytesIO
import re
from xml.sax.saxutils import escape
from fastapi import Response


def flat_rows(rows):
    def flatten(row, prefix=""):
        result = {}
        for key, value in row.items():
            label = (prefix + " " + key.replace("_", " ")).strip()
            if isinstance(value, dict):
                result.update(flatten(value, label))
            elif not isinstance(value, (list, tuple)):
                result[label] = value
        return result
    return [flatten(row) for row in rows]


def workbook_bytes(title, sections):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for name, source in sections.items():
        sheet = book.create_sheet(name[:31])
        rows = flat_rows(source)
        headers = list(dict.fromkeys(key for row in rows for key in row)) or ["Result"]
        sheet.append(headers)
        for row in rows:
            sheet.append([None] * len(headers))
            for column, key in enumerate(headers, 1):
                value = row.get(key)
                if isinstance(value, datetime):
                    value = value.isoformat()
                if isinstance(value, str):
                    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
                cell = sheet.cell(sheet.max_row, column)
                cell.value = value
                # User-authored names beginning '=' must never become formulas.
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if not rows:
            sheet.append(["No data in this period"])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(40, max(16, len(header) + 3))
    book.properties.title = title
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def pdf_bytes(title, sections, days):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, PageBreak
    output = BytesIO()
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("ReportCell", parent=styles["BodyText"], fontSize=8, leading=11, wordWrap="CJK")
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=32, bottomMargin=32)
    story = [Paragraph(escape(title), styles["Title"]), Paragraph(f"Last {days} days | Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC", styles["Normal"]), Spacer(1, 14)]
    story.extend([Paragraph("Font limitation: Hindi and emoji may display as missing glyphs in this PDF. Use Excel export to preserve the original Unicode text.", styles["Normal"]), Spacer(1, 12)])
    for name, source in sections.items():
        rows = flat_rows(source)
        if not rows:
            story.extend([Paragraph(escape(name), styles["Heading2"]), Paragraph("No data in this period.", styles["Normal"]), Spacer(1, 12)])
            continue
        headers = list(dict.fromkeys(key for row in rows for key in row))
        # Split wide reports into column groups, retaining the identifying column.
        groups = [headers[:7]] + [[headers[0]] + headers[i:i + 6] for i in range(7, len(headers), 6)]
        for group in groups:
            story.append(Paragraph(escape(name), styles["Heading2"]))
            data = [[Paragraph(escape(key.title()), cell_style) for key in group]]
            for row in rows:
                data.append([Paragraph(escape(str(row.get(key) if row.get(key) is not None else "—")), cell_style) for key in group])
            table = LongTable(data, colWidths=[doc.width / len(group)] * len(group), repeatRows=1, splitInRow=1)
            table.setStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                            ("TOPPADDING", (0, 0), (-1, -1), 7)])
            story.extend([table, Spacer(1, 16)])
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(30, 17, "SocialPilot | ROI uses entered budget and attributed revenue; missing provider metrics may be unavailable.")
        canvas.drawRightString(landscape(A4)[0] - 30, 17, str(document.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def response(kind, format, sections, days):
    title = "SocialPilot " + kind.replace("-", " ").title() + " Report"
    data = workbook_bytes(title, sections) if format == "xlsx" else pdf_bytes(title, sections, days)
    media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if format == "xlsx" else "application/pdf"
    return Response(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="socialpilot-{kind}.{format}"', "Cache-Control": "no-store"})
