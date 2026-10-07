"""Team isolation, approval gating, native exports and opt-in notification delivery."""
import unittest
from datetime import datetime, timedelta, timezone
from io import BytesIO
from unittest.mock import Mock, patch
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook
import models
import main
import notifications
from publishing_worker import run_due
import test_roles_api


class CollaborationTests(unittest.TestCase):
    def setUp(self):
        test_roles_api.RoleApiTests.setUp(self)
        self.admin.role = 'content_creator'
        self.db.commit()
    tearDown = test_roles_api.RoleApiTests.tearDown
    request = test_roles_api.RoleApiTests.request

    def team(self, approval=True):
        code, result = self.request('POST', '/teams', self.admin, {'name': 'Launch team', 'require_approval': approval})
        self.assertEqual(code, 201, result)
        self.assertEqual(self.request('PATCH', f"/teams/{result['id']}/members/{self.admin.id}", self.admin, {'role': 'content_creator'})[0], 200)
        return result['id']

    def join(self, team, user=None, role='content_creator'):
        user = user or self.user
        self.assertEqual(self.request('POST', f'/teams/{team}/invitations', self.admin, {'email': user.email, 'role': role})[0], 201)
        self.assertEqual(self.request('POST', f'/teams/{team}/accept', user)[0], 200)

    def draft(self, team, user=None, title='New launch'):
        code, result = self.request('POST', '/posts', user or self.user, {'title': title, 'content': 'Launch content'}, workspace=team)
        self.assertEqual(code, 201, result)
        return result

    def review(self, team, post, action, user=None, revision=1, comment=''):
        return self.request('POST', f'/posts/{post}/reviews', user or self.user,
                            {'action': action, 'revision': revision, 'comment': comment}, workspace=team)

    def test_workspace_and_personal_data_are_isolated_and_invites_require_acceptance(self):
        personal = self.draft(None, self.admin)
        team = self.team()
        self.assertEqual(self.request('GET', '/posts', self.user, workspace=team)[0], 403)
        self.request('POST', f'/teams/{team}/invitations', self.admin, {'email': self.user.email})
        self.assertEqual(self.request('GET', '/posts', self.user, workspace=team)[0], 403)
        self.request('POST', f'/teams/{team}/accept', self.user)
        post = self.draft(team)
        self.assertEqual([p['id'] for p in self.request('GET', '/posts', self.admin, workspace=team)[1]], [post['id']])
        self.assertEqual(self.request('GET', f"/posts/{personal['id']}", self.user, workspace=team)[0], 404)
        self.assertEqual(self.request('GET', f"/posts/{post['id']}", self.user)[0], 404)
        other = self.team()
        self.assertEqual(self.request('GET', f"/posts/{post['id']}", self.admin, workspace=other)[0], 404)

    def test_revocation_owner_protection_and_system_privilege_separation(self):
        team = self.team()
        self.join(team, role='administrator')
        self.request('PATCH', f'/teams/{team}/members/{self.user.id}', self.admin, {'is_admin': True})
        self.assertEqual(self.request('GET', '/admin/users', self.user, workspace=team)[0], 404)
        self.assertEqual(self.request('DELETE', f'/teams/{team}/members/{self.admin.id}', self.user)[0], 409)
        self.assertEqual(self.request('PATCH', f'/teams/{team}/members/{self.admin.id}', self.user, {'role': 'business_user'})[0], 409)
        self.request('PATCH', f'/teams/{team}/members/{self.user.id}', self.admin, {'role': 'business_user'})
        self.assertEqual(self.request('POST', '/posts', self.user, {'title': 'Forbidden'}, workspace=team)[0], 403)
        self.request('DELETE', f'/teams/{team}/members/{self.user.id}', self.admin)
        self.assertEqual(self.request('GET', '/reports/campaigns', self.user, workspace=team)[0], 403)
        self.assertEqual(self.request('GET', '/notifications', self.user)[1], [])

    def test_approval_revision_self_review_and_edit_invalidation(self):
        team = self.team()
        self.join(team, role='marketing_manager')
        post = self.draft(team)['id']
        self.assertEqual(self.review(team, post, 'approve', self.admin)[0], 409)
        self.assertEqual(self.review(team, post, 'submit')[0], 200)
        self.assertEqual(self.review(team, post, 'approve')[1]['approval_status'], 'awaiting_creator')
        self.assertEqual(self.review(team, post, 'submit')[0], 409)
        self.assertEqual(self.review(team, post, 'reject', self.admin)[0], 422)
        self.assertEqual(self.review(team, post, 'approve', self.admin, revision=2)[0], 409)
        self.assertEqual(self.review(team, post, 'approve', self.admin)[0], 200)
        code, result = self.request('PATCH', f'/posts/{post}', self.user, {'content': 'Changed'}, workspace=team)
        self.assertEqual(code, 200)
        self.assertEqual((result['revision'], result['approval_status']), (2, 'draft'))
        self.assertEqual(len(self.request('GET', f'/posts/{post}/reviews', self.user, workspace=team)[1]), 3)

    def test_schedule_gate_and_worker_defense(self):
        team = self.team()
        self.join(team, role='marketing_manager')
        post_id = self.draft(team)['id']
        owner = self.db.get(models.Workspace, team).storage_user_id
        account = models.SocialAccount(user_id=owner, platform='facebook', account_name='Page', connected=True)
        self.db.add(account); self.db.flush()
        self.db.add(models.SocialCredential(account_id=account.id, user_id=owner, provider='facebook', external_id='page', encrypted_token='unused'))
        self.db.commit()
        body = {'scheduled_at': (datetime.now(timezone.utc) + timedelta(hours=2)).isoformat(), 'account_ids': [account.id]}
        self.assertEqual(self.request('POST', f'/posts/{post_id}/schedule', self.admin, body, workspace=team)[0], 409)
        self.review(team, post_id, 'approve')
        self.review(team, post_id, 'approve', self.admin)
        code, jobs = self.request('POST', f'/posts/{post_id}/schedule', self.admin, body, workspace=team)
        self.assertEqual(code, 201, jobs)
        job = self.db.get(models.PublishingJob, jobs[0]['id'])
        job.run_at = datetime.now(timezone.utc) - timedelta(minutes=1)
        self.db.get(models.Post, post_id).approval_status = 'rejected'
        self.db.commit()
        publisher = Mock(return_value='unexpected')
        run_due(session_factory=sessionmaker(bind=self.engine), publisher=publisher)
        publisher.assert_not_called()
        self.db.expire_all()
        self.assertEqual(self.db.get(models.PublishingJob, job.id).status, 'failed')

    def test_workspace_principal_cannot_login_and_is_hidden(self):
        team = self.team()
        principal = self.db.get(models.User, self.db.get(models.Workspace, team).storage_user_id)
        principal.email = 'workspace-principal@example.com'
        self.db.commit()
        self.assertEqual(self.request('GET', '/me', principal)[0], 401)
        self.assertEqual(self.request('POST', '/login', body={'email': principal.email, 'password': 'anything'})[0], 401)
        self.assertEqual(self.request('GET', '/admin/users', self.admin)[0], 404)

    def test_inbox_read_state_and_preferences(self):
        notifications.notify_user(self.db, self.user.id, 'team', 'First update', key='unique')
        notifications.notify_user(self.db, self.user.id, 'team', 'Duplicate', key='unique')
        self.db.commit()
        self.assertEqual(self.db.query(models.NotificationDelivery).count(), 0)
        code, inbox = self.request('GET', '/notifications', self.user)
        self.assertEqual(code, 200)
        self.assertEqual(len(inbox), 1)
        self.assertFalse(inbox[0]['read'])
        self.assertEqual(self.request('POST', f"/notifications/{inbox[0]['id']}/read", self.admin)[0], 404)
        self.request('POST', '/notifications/read-all', self.user)
        self.assertEqual(self.request('GET', '/notifications', self.user, query='unread=true')[1], [])
        self.assertEqual(self.request('PUT', '/notifications/preferences', self.user, {'email': True})[0], 200)
        notifications.notify_user(self.db, self.user.id, 'team', 'Email update')
        self.db.commit()
        sender = Mock()
        with patch('notifications.email_ready', return_value=True):
            notifications.deliver_batch(sessionmaker(bind=self.engine), email_sender=sender)
            notifications.deliver_batch(sessionmaker(bind=self.engine), email_sender=sender)
        self.assertEqual(sender.call_count, 1)

    def test_disabled_delivery_and_revoked_membership_do_not_send(self):
        team = self.team(); self.join(team)
        self.request('PUT', '/notifications/preferences', self.user, {'email': True})
        notifications.notify_scope(self.db, self.db.get(models.Workspace, team).storage_user_id, 'team', 'Private team update')
        self.db.commit()
        self.request('DELETE', f'/teams/{team}/members/{self.user.id}', self.admin)
        sender = Mock()
        with patch('notifications.email_ready', return_value=True):
            notifications.deliver_batch(sessionmaker(bind=self.engine), email_sender=sender)
        sender.assert_not_called()

    def test_push_endpoint_rejects_arbitrary_network_targets(self):
        for endpoint in ['https://127.0.0.1/private', 'https://example.com/push', 'http://fcm.googleapis.com/push']:
            self.assertEqual(self.request('POST', '/notifications/subscriptions', self.user, {'endpoint': endpoint, 'keys': {}})[0], 422)

    def test_push_delivery_is_opt_in_and_failed_sends_are_not_duplicated(self):
        import base64
        encode = lambda value: base64.urlsafe_b64encode(value).decode().rstrip('=')
        body = {'endpoint': 'https://fcm.googleapis.com/fcm/send/test-only',
                'keys': {'p256dh': encode(b'\x04' + b'x' * 64), 'auth': encode(b'a' * 16)}}
        self.assertEqual(self.request('POST', '/notifications/subscriptions', self.user, body)[0], 201)
        self.assertEqual(self.request('POST', '/notifications/subscriptions', self.admin, body)[0], 409)
        self.request('PUT', '/notifications/preferences', self.user, {'push': True})
        notifications.notify_user(self.db, self.user.id, 'account', 'Push test')
        self.db.commit()
        sender = Mock(side_effect=RuntimeError('Sensitive provider error must stay private'))
        with patch('notifications.push_ready', return_value=True):
            notifications.deliver_batch(sessionmaker(bind=self.engine), push_sender=sender)
            notifications.deliver_batch(sessionmaker(bind=self.engine), push_sender=sender)
        self.assertEqual(sender.call_count, 1)
        result = self.request('GET', '/notifications/deliveries', self.user)[1]
        self.assertEqual(result[0]['status'], 'failed')
        self.assertNotIn('Sensitive', result[0]['error'])

    def test_periodic_reminders_and_campaign_alerts_are_deduplicated(self):
        now = datetime.now(timezone.utc)
        post = models.Post(user_id=self.user.id, title='Upcoming', content='Text', status='scheduled')
        self.db.add(post); self.db.flush()
        self.db.add(models.PublishingJob(post_id=post.id, user_id=self.user.id, provider='facebook', content='Text',
            scheduled_at=now + timedelta(hours=1), run_at=now + timedelta(hours=1), timezone_name='UTC', status='scheduled'))
        self.db.add(models.Campaign(user_id=self.user.id, name='Ending campaign', start_date=now.date() - timedelta(days=3),
            end_date=now.date(), budget=100, revenue=0, currency='INR'))
        self.db.commit()
        notifications.scan_events(self.db)
        first = self.db.query(models.Notification).count()
        notifications.scan_events(self.db)
        self.assertEqual(self.db.query(models.Notification).count(), first)
        self.assertEqual({n.category for n in self.db.query(models.Notification).all()}, {'reminder', 'campaign'})

    def test_report_export_contains_only_selected_workspace(self):
        team = self.team(False); self.join(team)
        owner = self.db.get(models.Workspace, team).storage_user_id
        today = datetime.now(timezone.utc).date()
        for user_id, name in [(owner, 'Team results'), (self.user.id, 'Private results')]:
            self.db.add(models.Campaign(user_id=user_id, creator_id=self.user.id, name=name, start_date=today, end_date=today,
                budget=100, revenue=200, currency='INR'))
        self.db.commit()
        code, data = self.request('GET', '/reports/export/campaigns', self.user, query='format=xlsx', workspace=team, raw=True)
        self.assertEqual(code, 200)
        book = load_workbook(BytesIO(data))
        values = str(list(book['Campaigns'].values))
        self.assertIn('Team results', values)
        self.assertNotIn('Private results', values)

    def test_native_reports_and_formula_safety(self):
        from report_exports import workbook_bytes, pdf_bytes
        book = load_workbook(BytesIO(workbook_bytes('Safety', {'Rows': [{'Name': '=HYPERLINK("evil")', 'Count': 42}]})))
        self.assertEqual(book.active['A2'].data_type, 's')
        self.assertEqual(book.active['B2'].value, 42)
        self.assertTrue(pdf_bytes('Test', {'Rows': [{'Name': '<b>Literal</b>', 'Count': 42}]}, 30).startswith(b'%PDF'))
        for kind in ['campaigns', 'audience', 'performance', 'engagement']:
            for format in ['xlsx', 'pdf']:
                code, data = self.request('GET', f'/reports/export/{kind}', self.user, query=f'format={format}&days=30', raw=True)
                self.assertEqual(code, 200, (kind, format, data[:200]))
                self.assertTrue(data.startswith(b'PK' if format == 'xlsx' else b'%PDF'))
        team = self.team()
        self.assertEqual(self.request('GET', '/reports/export/campaigns', self.user, workspace=team)[0], 403)
